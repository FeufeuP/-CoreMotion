"""Sistema visual neo-brutalista do CoreMotion.

Baseado nos tokens do tema:
  background #FFFFFF · foreground #000000 · primary #FF3333 · secondary #FFFF00
  accent #0066FF · chart-1..5 · radius 0 · sombra dura 4px · DM Sans + Space Mono
"""
from math import sqrt
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from lxml import etree

# ------------------------------------------------------------------ tokens
BG        = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x00, 0x00, 0x00)
INK_SOFT  = RGBColor(0x33, 0x33, 0x33)   # muted-foreground
MUTED     = RGBColor(0xF0, 0xF0, 0xF0)   # muted
BORDER    = RGBColor(0x00, 0x00, 0x00)
PRIMARY   = RGBColor(0xFF, 0x33, 0x33)   # chart-1 / primary
SECONDARY = RGBColor(0xFF, 0xFF, 0x00)   # chart-2 / secondary
ACCENT    = RGBColor(0x00, 0x66, 0xFF)   # chart-3 / accent
GREEN     = RGBColor(0x00, 0xCC, 0x00)   # chart-4
MAGENTA   = RGBColor(0xCC, 0x00, 0xCC)   # chart-5
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)

# dark variant
D_BG      = RGBColor(0x00, 0x00, 0x00)
D_CARD    = RGBColor(0x33, 0x33, 0x33)
D_PRIMARY = RGBColor(0xFF, 0x66, 0x66)
D_SECOND  = RGBColor(0xFF, 0xFF, 0x33)
D_ACCENT  = RGBColor(0x33, 0x99, 0xFF)
D_TEXT    = RGBColor(0xFF, 0xFF, 0xFF)
D_MUTED   = RGBColor(0xCC, 0xCC, 0xCC)

SANS = "DM Sans"
MONO = "Space Mono"

SHAPE = {
    "rect": MSO_SHAPE.RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
    "tri": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "star": MSO_SHAPE.STAR_8_POINT,
    "star4": MSO_SHAPE.STAR_4_POINT,
    "plus": MSO_SHAPE.MATH_PLUS,
    "diamond": MSO_SHAPE.DIAMOND,
    "hex": MSO_SHAPE.HEXAGON,
    "arrow": MSO_SHAPE.RIGHT_ARROW,
    "chev": MSO_SHAPE.CHEVRON,
}


# ------------------------------------------------------------------ sombra dura
def hard_shadow(shape, offset_px=6, color="000000", opacity=100):
    shape.shadow.inherit = False
    spPr = shape._element.spPr
    lst = spPr.find(qn("a:effectLst"))
    if lst is None:
        lst = etree.SubElement(spPr, qn("a:effectLst"))
    for child in list(lst):
        lst.remove(child)
    dist = int(round(sqrt(2) * offset_px * 9525))
    xml = (
        '<a:outerShdw xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        f'blurRad="0" dist="{dist}" dir="2700000" algn="ctr" rotWithShape="0">'
        f'<a:srgbClr val="{color}"><a:alpha val="{int(opacity * 1000)}"/></a:srgbClr>'
        "</a:outerShdw>"
    )
    lst.append(parse_xml(xml))


def _alpha(shape, pct, target="fill"):
    spPr = shape._element.spPr
    node = spPr.find(qn("a:solidFill")) if target == "fill" else None
    if target == "line":
        ln = spPr.find(qn("a:ln"))
        node = ln.find(qn("a:solidFill")) if ln is not None else None
    if node is None:
        return
    clr = node.find(qn("a:srgbClr"))
    if clr is None:
        return
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(pct * 1000)))


def box(slide, x, y, w, h, fill=WHITE, border=INK, lw=2.0, kind="rect",
        rot=0, shadow=True, shadow_px=6, alpha=None, dash=False,
        shadow_color="000000", radius=None):
    s = slide.shapes.add_shape(SHAPE.get(kind, MSO_SHAPE.RECTANGLE),
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        if alpha is not None:
            _alpha(s, alpha)
    if border is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = border
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = 4
    if rot:
        s.rotation = rot
    if shadow:
        hard_shadow(s, offset_px=shadow_px, color=shadow_color)
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


# ------------------------------------------------------------------ texto
def text(slide, x, y, w, h, items, size=13, color=INK, bold=False, font=SANS,
         spc=0, align="l", anchor="t", line=1.15, caps=False, italic=False,
         after=0, alpha=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    if isinstance(items, str):
        items = [items]
    flat = []
    for it in items:
        if isinstance(it, tuple):
            t, o = it
        else:
            t, o = it, {}
        for part in str(t).split("\n"):
            flat.append((part, dict(o)))

    for i, (t, o) in enumerate(flat):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                       "r": PP_ALIGN.RIGHT}[o.get("align", align)]
        p.line_spacing = o.get("line", line)
        p.space_after = Pt(o.get("after", after))
        p.space_before = Pt(o.get("before", 0))
        r = p.add_run()
        r.text = t.upper() if o.get("caps", caps) else t
        f = r.font
        f.name = o.get("font", font)
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.italic = o.get("italic", italic)
        f.color.rgb = o.get("color", color)
        sp = o.get("spc", spc)
        if sp:
            f._rPr.set("spc", str(int(sp * 100)))
        al = o.get("alpha", alpha)
        if al is not None:
            sf = f._rPr.find(qn("a:solidFill"))
            if sf is not None:
                clr = sf.find(qn("a:srgbClr"))
                if clr is not None:
                    a = etree.SubElement(clr, qn("a:alpha"))
                    a.set("val", str(int(al * 1000)))
    return tb


def mono(slide, x, y, w, h, txt, size=8.5, color=INK, bold=True, spc=2.0,
         align="l", anchor="t"):
    return text(slide, x, y, w, h, txt, size=size, color=color, bold=bold,
                font=MONO, spc=spc, align=align, anchor=anchor, caps=True)


def rule(slide, x, y, w, h=0.012, color=INK):
    return box(slide, x, y, w, h, fill=color, border=None, shadow=False)


def sticker(slide, x, y, w, h, label, fill=WHITE, color=INK, size=8,
           spc=1.6, shadow_px=5, border=2.0, dot=None):
    b = box(slide, x, y, w, h, fill=fill, border=INK, lw=border,
            shadow=True, shadow_px=shadow_px)
    if dot is not None:
        box(slide, x + 0.16, y + h / 2 - 0.045, 0.09, 0.09, fill=dot,
            border=INK, lw=1.2, kind="oval", shadow=False)
        mono(slide, x + 0.33, y, w - 0.40, h, label, size=size, color=color,
             spc=spc, anchor="m")
    else:
        mono(slide, x, y, w, h, label, size=size, color=color, spc=spc,
             align="c", anchor="m")
    return b


# ------------------------------------------------------------------ chrome
TOTAL = 16
PROGRESS_COLORS = [PRIMARY, SECONDARY, ACCENT, GREEN, MAGENTA]


def chrome(slide, idx, section, dark=False, total=TOTAL, accent=PRIMARY):
    """Cabeçalho mínimo: marca, seção, contador e fio inferior."""
    fg = D_TEXT if dark else INK
    soft = D_MUTED if dark else INK_SOFT
    name_it(box(slide, 0.78, 0.40, 0.15, 0.15, fill=accent, border=None,
                shadow=False), "CM!!dot")
    name_it(mono(slide, 1.04, 0.38, 3.4, 0.2, "CoreMotion", size=9.5, color=fg,
                 spc=2.6), "CM!!brand")
    name_it(mono(slide, 7.30, 0.38, 5.25, 0.2,
                 f"{section}   ·   {idx:02d}/{total}", size=8, color=soft,
                 spc=2.0, align="r"), "CM!!count")
    name_it(rule(slide, 0, 7.44, 13.333, 0.06, color=fg), "CM!!rule")


# ------------------------------------------------------------------ título
def title_block(slide, kicker, headline, sub=None, y=1.12, size=36,
                kicker_color=PRIMARY, sub_w=9.4, dark=False):
    fg = D_TEXT if dark else INK
    soft = D_MUTED if dark else INK_SOFT
    k = mono(slide, 0.78, y, 8.0, 0.22, "// " + kicker, size=8.5,
             color=kicker_color, spc=2.4)
    h = text(slide, 0.78, y + 0.30, 11.6, 0.9, headline, size=size,
             color=fg, bold=True, spc=-1.1, line=0.98)
    name_it(k, "CM!!kick")
    name_it(h, "CM!!head")
    out = [k, h]
    if sub:
        out.append(text(slide, 0.78, y + 1.14, sub_w, 0.5, sub, size=12.5,
                        color=soft, line=1.3))
    return out


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def cascade(shapes, start=0, step=80, effect="rise", dur=520):
    return [(s, effect, start + i * step, dur) for i, s in enumerate(shapes)]


# ------------------------------------------------------------------ kit velocidade / esporte
def name_it(shape, name):
    """Dá um nome ao shape — nomes iguais com '!!' forçam morph entre slides."""
    try:
        shape.name = name
    except Exception:
        try:
            shape._element.nvSpPr.cNvPr.set("name", name)
        except Exception:
            pass
    return shape


def send_to_back(slide, shapes):
    """Move shapes para o fundo do slide (acima do fundo, atrás do conteúdo)."""
    if not isinstance(shapes, (list, tuple)):
        shapes = [shapes]
    spTree = slide._element.find(qn("p:cSld")).find(qn("p:spTree"))
    anchor = 0
    for i, child in enumerate(spTree):
        tag = child.tag
        if tag.endswith("}nvGrpSpPr") or tag.endswith("}grpSpPr"):
            anchor = i + 1
    for sh in reversed(list(shapes)):
        el = sh._element
        spTree.remove(el)
        spTree.insert(anchor, el)
    return shapes


def speed_streaks(slide, x, y, w, n=3, gap=0.15, color=INK, alpha=16,
                  thick=0.055, taper=0.62, rot=0):
    """Rastros de movimento: linhas horizontais que encurtam, como blur de velocidade."""
    out = []
    for i in range(n):
        ww = w * (taper ** i)
        out.append(box(slide, x, y + i * gap, ww, thick, fill=color,
                       border=None, shadow=False, alpha=alpha, rot=rot))
    return out


def chevrons(slide, x, y, n=3, size=0.26, gap=0.34, color=PRIMARY, lw=2.0,
             rot=0, last_solid=True):
    """Fileira de chevrons '>>>' apontando para frente (ritmo, avanço, velocidade)."""
    out = []
    for i in range(n):
        solid = last_solid and i == n - 1
        out.append(box(slide, x + i * gap, y, size * 0.85, size,
                       fill=(color if solid else None),
                       border=(None if solid else color), lw=lw,
                       kind="chev", shadow=False, rot=rot))
    return out


def checkered(slide, x, y, cols=6, rows=2, cell=0.16, c1=INK, c2=WHITE,
              border=None, lw=1.0, shadow=False):
    """Bandeira quadriculada de largada/chegada (F1, atletismo)."""
    out = []
    for r in range(rows):
        for c in range(cols):
            col = c1 if (r + c) % 2 == 0 else c2
            out.append(box(slide, x + c * cell, y + r * cell, cell, cell,
                           fill=col, border=border, lw=lw, shadow=shadow))
    return out


def speedometer(slide, x, y, size=1.25, ring=PRIMARY, bg=WHITE, needle=INK,
                hub=None, angle=38, lw=2.2):
    """Velocímetro semicircular: anel + agulha + cubo. bg deve casar com o fundo."""
    parts = []
    parts.append(box(slide, x, y, size, size, fill=ring, border=INK, lw=lw,
                     kind="oval", shadow=True, shadow_px=5))
    inner = size * 0.62
    parts.append(box(slide, x + (size - inner) / 2, y + (size - inner) / 2,
                     inner, inner, fill=bg, border=INK, lw=1.8, kind="oval",
                     shadow=False))
    # corta a metade inferior para virar mostrador
    parts.append(box(slide, x - 0.03, y + size * 0.52, size + 0.06,
                     size * 0.62, fill=bg, border=None, shadow=False))
    # agulha com pivô no centro do mostrador
    from math import sin, cos, radians
    ln = size * 0.34
    th = max(0.05, size * 0.075)
    a = radians(angle)
    cx, cy = x + size / 2, y + size * 0.52
    mx = cx + (ln / 2) * sin(a)
    my = cy - (ln / 2) * cos(a)
    parts.append(box(slide, mx - th / 2, my - ln / 2, th, ln, fill=needle,
                     border=None, shadow=False, rot=angle))
    parts.append(box(slide, cx - size * 0.075, cy - size * 0.075,
                     size * 0.15, size * 0.15, fill=(hub or needle),
                     border=INK, lw=1.4, kind="oval", shadow=False))
    return parts


def ghost(slide, x, y, w, h, word, size=90, color=INK, alpha=6, font=SANS,
          align="l", rot=0):
    """Palavra gigante quase invisível no fundo — cartaz de estádio."""
    g = text(slide, x, y, w, h, word, size=size, color=color, bold=True,
             font=font, align=align, line=0.9, alpha=alpha)
    if rot:
        g.rotation = rot
    return g
