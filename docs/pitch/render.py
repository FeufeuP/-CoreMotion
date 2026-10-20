"""Renderiza o deck .pptx em PNGs (aproximação fiel) para conferência visual."""
import sys
import re
import os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Length, Emu

SRC = sys.argv[1] if len(sys.argv) > 1 else "../CoreMotion-Pitch-Banca.pptx"
OUTDIR = sys.argv[2] if len(sys.argv) > 2 else "/tmp/preview"
DPI = int(sys.argv[3]) if len(sys.argv) > 3 else 130

REG = "/tmp/arimo/Arimo-Regular.ttf"
BOLD = "/tmp/arimo/Arimo-Bold.ttf"
ITAL = "/tmp/arimo/Arimo-Italic.ttf"

FAMILIES = {
    "DM Sans": ("/tmp/fonts/DMSans-Regular.ttf", "/tmp/fonts/DMSans-Bold.ttf",
                "/tmp/fonts/DMSans-Regular.ttf"),
    "Space Mono": ("/tmp/fonts/SpaceMono-Regular.ttf", "/tmp/fonts/SpaceMono-Bold.ttf",
                   "/tmp/fonts/SpaceMono-Regular.ttf"),
}
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

prs = Presentation(SRC)
S = DPI / 914400.0
W = int(round(prs.slide_width * S))
H = int(round(prs.slide_height * S))
os.makedirs(OUTDIR, exist_ok=True)
_cache = {}


def fnt(size_pt, bold=False, italic=False, family=None):
    key = (round(size_pt, 1), bold, italic, family)
    if key not in _cache:
        reg, bold_f, ital_f = FAMILIES.get(family, (REG, BOLD, ITAL))
        path = bold_f if bold else (ital_f if italic else reg)
        _cache[key] = ImageFont.truetype(path, max(5, int(round(size_pt / 72 * DPI))))
    return _cache[key]


def alpha_of(clr):
    if clr is None:
        return 255
    a = clr.find(A + "alpha")
    return 255 if a is None else int(int(a.get("val")) / 100000 * 255)


def hexrgb(clr):
    v = clr.get("val")
    return (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))


def fill_of(shape):
    try:
        node = shape._element.spPr.find(A + "solidFill")
    except Exception:
        return None
    if node is None:
        return None
    clr = node.find(A + "srgbClr")
    if clr is None:
        return None
    return hexrgb(clr) + (alpha_of(clr),)


def line_of(shape):
    try:
        ln = shape._element.spPr.find(A + "ln")
    except Exception:
        return None, 0, False
    if ln is None:
        return None, 0, False
    dash = ln.find(A + "prstDash") is not None
    w = max(1, int(round(int(ln.get("w") or 12700) / 914400 * DPI)))
    node = ln.find(A + "solidFill")
    if node is None:
        return None, w, dash
    clr = node.find(A + "srgbClr")
    if clr is None:
        return None, w, dash
    return hexrgb(clr) + (alpha_of(clr),), w, dash


def run_props(r):
    size = r.font.size.pt if r.font.size else 14.0
    color, alpha = (255, 255, 255), 255
    try:
        sf = r.font._rPr.find(A + "solidFill")
        if sf is not None:
            clr = sf.find(A + "srgbClr")
            if clr is not None:
                color = hexrgb(clr)
                alpha = alpha_of(clr)
    except Exception:
        pass
    spc = 0.0
    if r.font._rPr is not None and r.font._rPr.get("spc"):
        spc = float(r.font._rPr.get("spc")) / 100 / 72 * DPI
    return {"text": r.text,
            "f": fnt(size, bool(r.font.bold), bool(r.font.italic), r.font.name),
            "color": color, "alpha": alpha, "spc": spc, "size": size}


def text_width(s, run):
    return run["f"].getlength(s) + run["spc"] * len(s)


def draw_text(d, shape):
    tf = shape.text_frame
    x0, y0 = shape.left * S, shape.top * S
    w, h = shape.width * S, shape.height * S
    lines = []
    for p in tf.paragraphs:
        runs = [run_props(r) for r in p.runs if r.text]
        if not runs:
            continue
        base = max(r["size"] for r in runs)
        ls = p.line_spacing
        if isinstance(ls, Length):
            lh = ls.pt / 72 * DPI
        elif isinstance(ls, (int, float)):
            lh = base * 1.2 * float(ls) / 72 * DPI
        else:
            lh = base * 1.2 / 72 * DPI
        para = {"lh": lh, "after": (p.space_after.pt if p.space_after else 0) / 72 * DPI,
                "before": (p.space_before.pt if p.space_before else 0) / 72 * DPI,
                "align": p.alignment or PP_ALIGN.LEFT}
        cur, cur_w = [], 0.0
        for run in runs:
            for tok in re.split(r"(\s+)", run["text"]):
                if tok == "":
                    continue
                tw = text_width(tok, run)
                if tok.strip() == "":
                    if cur:
                        cur.append((tok, run, tw))
                        cur_w += tw
                    continue
                if cur_w + tw > w + 0.5 and cur:
                    while cur and cur[-1][0].strip() == "":
                        cur_w -= cur[-1][2]
                        cur.pop()
                    lines.append((para, cur, cur_w))
                    cur, cur_w = [], 0.0
                cur.append((tok, run, tw))
                cur_w += tw
        if cur:
            while cur and cur[-1][0].strip() == "":
                cur_w -= cur[-1][2]
                cur.pop()
            lines.append((para, cur, cur_w))

    total = 0.0
    for i, (para, _, _) in enumerate(lines):
        total += para["before"] + para["lh"]
        if i == len(lines) - 1 or lines[i + 1][0] is not para:
            total += para["after"]

    cy = y0
    if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
        cy = y0 + (h - total) / 2
    elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
        cy = y0 + h - total

    for i, (para, toks, lw) in enumerate(lines):
        cy += para["before"]
        cx = x0
        if para["align"] == PP_ALIGN.RIGHT:
            cx = x0 + w - lw
        elif para["align"] == PP_ALIGN.CENTER:
            cx = x0 + (w - lw) / 2
        asc, _ = toks[0][1]["f"].getmetrics()
        base_y = cy + para["lh"] / 2 + asc / 2 - 0.05 * DPI
        for tok, run, tw in toks:
            if run["spc"]:
                for ch in tok:
                    d.text((cx, base_y), ch, font=run["f"],
                           fill=run["color"] + (run["alpha"],), anchor="ls")
                    cx += run["f"].getlength(ch) + run["spc"]
            else:
                d.text((cx, base_y), tok, font=run["f"],
                       fill=run["color"] + (run["alpha"],), anchor="ls")
                cx += tw
        cy += para["lh"]
        if i == len(lines) - 1 or lines[i + 1][0] is not para:
            cy += para["after"]


def prst_of(shape):
    try:
        return shape._element.spPr.prstGeom.get("prst")
    except Exception:
        return None


def draw_shape(d, shape, img):
    x, y = shape.left * S, shape.top * S
    w, h = shape.width * S, shape.height * S
    prst = prst_of(shape)
    f = fill_of(shape)
    l, lw, dash = line_of(shape)
    rot = shape.rotation or 0

    if rot:
        pad = 6
        tmp = Image.new("RGBA", (int(w) + pad * 2, int(h) + pad * 2), (0, 0, 0, 0))
        td = ImageDraw.Draw(tmp)
        shp = shape._element
        shdw = shp.spPr.find(A + "effectLst")
        if shdw is not None:
            off = shdw.find(".//" + A + "outerShdw")
            if off is not None:
                dist = int(off.get("dist") or 0) / 914400 * DPI
                td.rectangle([pad + dist * .7, pad + dist * .7,
                              pad + w + dist * .7, pad + h + dist * .7],
                             fill=(0, 0, 0, 255))
        _prim(td, prst, [pad, pad, pad + w, pad + h], f, l, lw, dash)
        tmp = tmp.rotate(-rot, resample=Image.BICUBIC, expand=True)
        ox = int(x + w / 2 - tmp.width / 2)
        oy = int(y + h / 2 - tmp.height / 2)
        img.paste(tmp, (ox, oy), tmp)
        return
    _prim(d, prst, [x, y, x + w, y + h], f, l, lw, dash)


def _prim(d, prst, box, f, l, lw, dash):
    if prst == "ellipse":
        d.ellipse(box, fill=f, outline=l, width=lw)
    elif prst == "roundRect":
        r = min(box[2] - box[0], box[3] - box[1]) * 0.16
        d.rounded_rectangle(box, radius=r, fill=f, outline=l, width=lw)
    elif prst == "triangle":
        d.polygon([((box[0] + box[2]) / 2, box[1]), (box[2], box[3]),
                   (box[0], box[3])], fill=f, outline=l)
    elif prst == "hexagon":
        w, h = box[2] - box[0], box[3] - box[1]
        d.polygon([(box[0] + w * .25, box[1]), (box[0] + w * .75, box[1]),
                   (box[2], box[1] + h / 2), (box[0] + w * .75, box[3]),
                   (box[0] + w * .25, box[3]), (box[0], box[1] + h / 2)],
                  fill=f, outline=l)
    elif prst == "diamond":
        d.polygon([((box[0] + box[2]) / 2, box[1]), (box[2], (box[1] + box[3]) / 2),
                   ((box[0] + box[2]) / 2, box[3]), (box[0], (box[1] + box[3]) / 2)],
                  fill=f, outline=l)
    elif prst == "chevron":
        w, h = box[2] - box[0], box[3] - box[1]
        d.polygon([(box[0], box[1]), (box[0] + w * .55, box[1]), (box[2], box[1] + h / 2),
                   (box[0] + w * .55, box[3]), (box[0], box[3]),
                   (box[0] + w * .45, box[1] + h / 2)], fill=f, outline=l)
    elif prst == "rtTriangle":
        d.polygon([(box[0], box[1]), (box[0], box[3]), (box[2], box[3])], fill=f)
    else:
        d.rectangle(box, fill=f, outline=l, width=lw)


def draw_connector(d, shape):
    x1, y1 = shape.begin_x * S, shape.begin_y * S
    x2, y2 = shape.end_x * S, shape.end_y * S
    l, lw, dash = line_of(shape)
    if l is None:
        return
    if dash:
        # tracejado manual
        from math import hypot
        total = hypot(x2 - x1, y2 - y1)
        step = 6
        t = 0
        while t < total:
            a = t / total
            b = min(1, (t + step) / total)
            d.line([x1 + (x2 - x1) * a, y1 + (y2 - y1) * a,
                    x1 + (x2 - x1) * b, y1 + (y2 - y1) * b], fill=l, width=lw)
            t += step * 2
    else:
        d.line([x1, y1, x2, y2], fill=l, width=lw)


files = []
for idx, slide in enumerate(prs.slides, start=1):
    bgc = (251, 247, 242)
    try:
        bgel = slide._element.find(P + "cSld").find(P + "bg")
        if bgel is not None:
            sf = bgel.find(".//" + A + "srgbClr")
            if sf is not None:
                bgc = hexrgb(sf)
    except Exception:
        pass
    img = Image.new("RGB", (W, H), bgc)
    canvas = ImageDraw.Draw(img, "RGBA")
    # ordem visual: shapes na ordem de inserção
    for shape in slide.shapes:
        st = shape.shape_type
        if shape._element.tag.endswith("cxnSp"):
            draw_connector(canvas, shape)
        elif st == MSO_SHAPE_TYPE.AUTO_SHAPE:
            draw_shape(canvas, shape, img)
            if shape.has_text_frame and shape.text_frame.text.strip():
                draw_text(canvas, shape)
        elif shape.has_text_frame and shape.text_frame.text.strip():
            draw_text(canvas, shape)
    out = os.path.join(OUTDIR, f"slide-{idx:02d}.png")
    img.save(out)
    files.append(out)

print("ok", len(files), "slides ->", OUTDIR)
